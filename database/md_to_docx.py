# -*- coding: utf-8 -*-
"""将 Markdown 文档转换为格式化的 Word(.docx) 文档（支持标题/段落/列表/表格/引用）"""
import re
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

HEI = "黑体"
SONG = "宋体"


def set_font(run, name=SONG, size=10.5, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color


def strip_inline(text):
    """去掉行内 markdown 标记（**加粗**、`代码`、[链接](url)）"""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text


def add_table(doc, lines):
    rows = []
    for ln in lines:
        if ln.strip().startswith("|"):
            cells = [c.strip() for c in ln.strip().strip("|").split("|")]
            rows.append(cells)
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    table = doc.add_table(rows=1, cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j in range(ncol):
        cell = table.rows[0].cells[j]
        cell.text = ""
        r = cell.paragraphs[0].add_run(strip_inline(rows[0][j]) if j < len(rows[0]) else "")
        set_font(r, HEI, 9.5, bold=True)
    for i in range(1, len(rows)):
        row = table.add_row()
        for j in range(ncol):
            cell = row.cells[j]
            cell.text = ""
            txt = strip_inline(rows[i][j]) if j < len(rows[i]) else ""
            r = cell.paragraphs[0].add_run(txt)
            set_font(r, SONG, 9)
    doc.add_paragraph()


def convert(md_path, docx_path):
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)

    i = 0
    first_h1_done = False
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        # 表格块
        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = strip_inline(m.group(2))
            if level == 1 and not first_h1_done:
                first_h1_done = True
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(text)
                set_font(r, HEI, 24, bold=True)
            else:
                h = doc.add_heading(level=min(level, 3))
                r = h.add_run(text)
                set_font(r, HEI, {1: 16, 2: 13, 3: 11.5}.get(min(level, 3), 11), bold=True)
            i += 1
            continue

        # 水平分割线
        if stripped == "---" or stripped == "***":
            i += 1
            continue

        # 引用
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            r = p.add_run(strip_inline(stripped.lstrip("> ")))
            set_font(r, SONG, 10, italic=True, color=RGBColor(0x55, 0x55, 0x55))
            i += 1
            continue

        # 无序列表
        if re.match(r"^\s*[-*+]\s+", stripped):
            m2 = re.match(r"^\s*[-*+]\s+(.*)$", stripped)
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(strip_inline(m2.group(1)))
            set_font(r, SONG, 10.5)
            i += 1
            continue

        # 有序列表
        m3 = re.match(r"^\s*(\d+)[.、]\s+(.*)$", stripped)
        if m3:
            p = doc.add_paragraph(style="List Number")
            r = p.add_run(strip_inline(m3.group(2)))
            set_font(r, SONG, 10.5)
            i += 1
            continue

        # 空行
        if stripped == "":
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        r = p.add_run(strip_inline(stripped))
        set_font(r, SONG, 10.5)
        i += 1

    doc.save(docx_path)
    print(f"[OK] 已生成 Word：{docx_path}")


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
