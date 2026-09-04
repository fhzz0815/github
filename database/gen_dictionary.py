# -*- coding: utf-8 -*-
"""
从 smart_restaurant.sql 解析表结构，生成数据库字典（V3.0 版）：
  - 第一部分：数据表基础说明（表名、中文释义、业务用途）
  - 第二部分：表间关联关系（ER，逻辑外键：父表/子表/关联字段/关联类型/级联策略）
  - 第三部分：字段详细说明（字段名、数据类型、约束、默认值、中文含义）
  - 附录：报表视图
输出：Markdown + Word
"""
import re
import sqlglot
from sqlglot import exp

SQL_PATH = r"D:\TraeProject\Smart Restaurant\database\smart_restaurant.sql"
MD_PATH  = r"D:\TraeProject\Smart Restaurant\docs\智慧餐厅_数据库字典.md"
DOCX_PATH = r"D:\TraeProject\Smart Restaurant\docs\智慧餐厅_数据库字典.docx"

with open(SQL_PATH, encoding="utf-8") as f:
    sql = f.read()
statements = sqlglot.parse(sql, read="mysql")


def clean_comment(v):
    if v is None:
        return ""
    if isinstance(v, exp.Expression):
        v = v.name if hasattr(v, "name") else v.sql()
    return str(v).strip("'\"")


def extract_str(expr):
    if expr is None:
        return None
    if isinstance(expr, exp.Literal):
        return expr.name
    if isinstance(expr, exp.Identifier):
        return expr.name
    return expr.sql()


def norm_type(t):
    if not t:
        return t
    t = re.sub(r"\s*,\s*", ",", t)
    t = t.replace("UBIGINT", "BIGINT UNSIGNED")
    t = t.replace("UTINYINT", "TINYINT UNSIGNED")
    t = t.replace("USMALLINT", "SMALLINT UNSIGNED")
    t = t.replace("UINT", "INT UNSIGNED")
    t = t.replace("UDECIMAL", "DECIMAL UNSIGNED")
    return t


def render_default(d):
    if d is None:
        return ""
    if isinstance(d, exp.Expression):
        return d.sql().replace("CURRENT_TIMESTAMP()", "CURRENT_TIMESTAMP")
    return clean_comment(d)


def parse_column(col):
    name = col.name
    kind = col.args.get("kind")
    data_type = kind.sql() if kind else ""
    nullable = True
    default = ""
    comment = ""
    is_pk = False
    is_auto = False
    is_unique = False
    for c in col.args.get("constraints", []) or []:
        k = c.kind if hasattr(c, "kind") else None
        if isinstance(k, exp.NotNullColumnConstraint):
            nullable = False
        elif isinstance(k, exp.DefaultColumnConstraint):
            d = k.args.get("this")
            if d is not None:
                default = render_default(d)
        elif isinstance(k, exp.CommentColumnConstraint):
            comment = clean_comment(k.args.get("this"))
        elif isinstance(k, exp.PrimaryKeyColumnConstraint):
            is_pk = True
            nullable = False
        elif isinstance(k, exp.AutoIncrementColumnConstraint):
            is_auto = True
        elif isinstance(k, exp.UniqueColumnConstraint):
            is_unique = True
    return {
        "name": name, "type": norm_type(data_type), "nullable": nullable,
        "default": default, "comment": comment, "is_pk": is_pk,
        "is_auto": is_auto, "is_unique": is_unique,
    }


def parse_create(stmt):
    schema = stmt.this
    tname = schema.this.name
    table_comment = ""
    props = stmt.args.get("properties")
    if props:
        for p in props.expressions:
            if isinstance(p, exp.SchemaCommentProperty):
                table_comment = clean_comment(p.args.get("this"))
    columns = []
    indexes = []
    for e in schema.expressions:
        if isinstance(e, exp.ColumnDef):
            columns.append(parse_column(e))
        elif isinstance(e, exp.PrimaryKey):
            cols = [c.name for c in e.expressions if isinstance(c, exp.Identifier)]
            indexes.append({"name": "PRIMARY", "unique": True, "cols": cols, "type": "PRIMARY"})
        elif isinstance(e, exp.UniqueColumnConstraint):
            sch = e.args.get("this")
            if isinstance(sch, exp.Schema):
                name = sch.this.name if hasattr(sch.this, "name") else ""
                cols = [c.name for c in sch.expressions if isinstance(c, exp.Identifier)]
            else:
                name, cols = "", []
            indexes.append({"name": name, "unique": True, "cols": cols, "type": "UNIQUE"})
        elif isinstance(e, exp.IndexColumnConstraint):
            nm = e.args.get("this")
            name = nm.name if isinstance(nm, exp.Identifier) else ""
            cols = []
            for o in e.expressions:
                col = o.this if isinstance(o, exp.Ordered) else o
                if isinstance(col, exp.Column):
                    cols.append(col.name)
                elif isinstance(col, exp.Identifier):
                    cols.append(col.name)
            indexes.append({"name": name, "unique": False, "cols": cols, "type": "INDEX"})
    return {"name": tname, "comment": table_comment, "columns": columns, "indexes": indexes}


tables = []
for st in statements:
    if isinstance(st, exp.Create) and isinstance(st.this, exp.Schema):
        tables.append(parse_create(st))
by_name = {t["name"]: t for t in tables}

# ---------------- 表业务用途清单（按 SQL 建表顺序） ----------------
TABLE_INFO = {
    "sys_role": "定义系统角色（管理员/店长/服务员/收银/后厨），支撑权限分配",
    "sys_permission": "定义商户端菜单与操作权限点，与角色多对多绑定",
    "sys_role_permission": "角色与权限的多对多中间表，实现 RBAC 权限分配",
    "sys_user": "员工/商户后台登录账号，含角色与所属门店",
    "staff_schedule": "员工排班计划（早晚班/时间），支撑门店人力管理",
    "staff_login_log": "员工登录日志，用于账号安全审计与异常登录排查",
    "store": "门店主数据，含地址/坐标/营业时间/配送参数",
    "store_payment_setting": "门店支付渠道配置（收款码/费率）",
    "member_category": "会员等级类别（银卡/金卡/黑金卡），定义充值门槛与折扣",
    "member": "会员主档案（微信openid/手机号/余额/积分/等级）",
    "member_address": "会员收货地址，支撑外卖配送",
    "member_bank_card": "会员绑定银行卡，支撑提现/退款原路返回",
    "member_balance_record": "会员钱包余额流水，记账可追溯、可对账",
    "member_points_record": "会员积分流水（获得/消费），支撑积分体系",
    "member_recharge_record": "会员充值记录（含赠送金额与支付流水）",
    "red_packet": "会员红包（充值奖励/营销发放）",
    "dish_category": "菜品分类（招牌热菜/家常小炒/饮品等）",
    "dish_taste": "菜品口味（辣度/做法等可选项）",
    "dish_spec": "菜品规格（大/中/小份等）",
    "ingredient_category": "原料类别（生鲜/粮油/调料等）",
    "ingredient": "原料档案（名称/规格/单位/参考成本）",
    "dish": "菜品主数据（价格/分类/口味/规格/上下架/销量）",
    "dish_ingredient_rel": "菜品与原料多对多关联，支撑成本核算与反查",
    "dish_review": "菜品评价（评分/图文/匿名/商家回复）",
    "table_type": "桌型定义（2人/4人/6人/包间）",
    "dining_table": "台桌档案（桌号/桌型/二维码/状态）",
    "queue": "排队取号记录（叫号/过号/就位），支撑取号叫号流程",
    "reservation": "餐桌预约（时段/人数/预留桌）",
    "cart": "购物车（会员会话级点餐暂存）",
    "orders": "订单主表（堂食/外卖，金额拆分，状态机，乐观锁）",
    "order_detail": "订单明细（菜品快照/数量/小计/制作状态/已退金额）",
    "order_status_log": "订单状态流转日志，业务审计留痕",
    "payment_record": "支付记录（渠道/幂等键/交易号/回调）",
    "refund": "退单/退款记录（退菜/整单，金额与原因）",
    "coupon": "优惠券定义（满减/折扣/立减，使用范围）",
    "member_coupon": "会员领用优惠券（状态/有效期/使用订单）",
    "dish_stock": "菜品库存（按门店按菜品）",
    "ingredient_stock": "原料库存（按门店按原料）",
    "stock_check_dish": "菜品盘点单（盘点批次/差异）",
    "stock_check_ingredient": "原料盘点单（盘点批次/差异）",
    "printer": "门店打印机绑定（打印任务/启用状态）",
    "receipt_template": "小票模板（结账单/外卖单/排队单模板内容）",
    "feedback": "客户意见反馈",
}

# ---------------- 逻辑外键（ER）关系清单 ----------------
# 格式: (子表, 父表, 子表关联字段, 父表关联字段, 级联策略)
# CASCADE=级联删除 RESTRICT=限制删除 LOGICAL=仅逻辑关联无约束
ER_DATA = [
    ("sys_role_permission", "sys_role", "role_id", "id", "CASCADE"),
    ("sys_role_permission", "sys_permission", "permission_id", "id", "CASCADE"),
    ("sys_user", "sys_role", "role_id", "id", "RESTRICT"),
    ("staff_schedule", "sys_user", "staff_id", "id", "CASCADE"),
    ("staff_login_log", "sys_user", "staff_id", "id", "CASCADE"),
    ("store_payment_setting", "store", "store_id", "id", "CASCADE"),
    ("member", "member_category", "member_category_id", "id", "RESTRICT"),
    ("member", "store", "register_store_id", "id", "RESTRICT"),
    ("member_address", "member", "member_id", "id", "CASCADE"),
    ("member_bank_card", "member", "member_id", "id", "CASCADE"),
    ("member_balance_record", "member", "member_id", "id", "CASCADE"),
    ("member_points_record", "member", "member_id", "id", "CASCADE"),
    ("member_recharge_record", "member", "member_id", "id", "RESTRICT"),
    ("member_recharge_record", "store", "store_id", "id", "RESTRICT"),
    ("red_packet", "member", "member_id", "id", "CASCADE"),
    ("dish_category", "store", "store_id", "id", "CASCADE"),
    ("dish_taste", "store", "store_id", "id", "CASCADE"),
    ("dish_spec", "store", "store_id", "id", "CASCADE"),
    ("ingredient_category", "store", "store_id", "id", "CASCADE"),
    ("ingredient", "store", "store_id", "id", "CASCADE"),
    ("ingredient", "ingredient_category", "category_id", "id", "RESTRICT"),
    ("dish", "store", "store_id", "id", "CASCADE"),
    ("dish", "dish_category", "category_id", "id", "RESTRICT"),
    ("dish", "dish_taste", "taste_id", "id", "RESTRICT"),
    ("dish", "dish_spec", "spec_id", "id", "RESTRICT"),
    ("dish_ingredient_rel", "dish", "dish_id", "id", "CASCADE"),
    ("dish_ingredient_rel", "ingredient", "ingredient_id", "id", "CASCADE"),
    ("dish_review", "dish", "dish_id", "id", "CASCADE"),
    ("dish_review", "member", "member_id", "id", "RESTRICT"),
    ("table_type", "store", "store_id", "id", "CASCADE"),
    ("dining_table", "store", "store_id", "id", "CASCADE"),
    ("dining_table", "table_type", "table_type_id", "id", "RESTRICT"),
    ("queue", "store", "store_id", "id", "CASCADE"),
    ("queue", "table_type", "table_type_id", "id", "RESTRICT"),
    ("reservation", "store", "store_id", "id", "CASCADE"),
    ("reservation", "member", "member_id", "id", "RESTRICT"),
    ("reservation", "dining_table", "table_id", "id", "RESTRICT"),
    ("cart", "member", "member_id", "id", "CASCADE"),
    ("cart", "dish", "dish_id", "id", "RESTRICT"),
    ("orders", "store", "store_id", "id", "RESTRICT"),
    ("orders", "member", "member_id", "id", "RESTRICT"),
    ("orders", "dining_table", "table_id", "id", "RESTRICT"),
    ("orders", "coupon", "coupon_id", "id", "LOGICAL"),
    ("order_detail", "orders", "order_id", "id", "CASCADE"),
    ("order_detail", "dish", "dish_id", "id", "RESTRICT"),
    ("order_status_log", "orders", "order_id", "id", "CASCADE"),
    ("payment_record", "orders", "order_id", "id", "RESTRICT"),
    ("payment_record", "store", "store_id", "id", "RESTRICT"),
    ("payment_record", "member", "member_id", "id", "RESTRICT"),
    ("refund", "orders", "order_id", "id", "RESTRICT"),
    ("coupon", "store", "store_id", "id", "RESTRICT"),
    ("member_coupon", "member", "member_id", "id", "CASCADE"),
    ("member_coupon", "coupon", "coupon_id", "id", "CASCADE"),
    ("dish_stock", "store", "store_id", "id", "CASCADE"),
    ("dish_stock", "dish", "dish_id", "id", "CASCADE"),
    ("ingredient_stock", "store", "store_id", "id", "CASCADE"),
    ("ingredient_stock", "ingredient", "ingredient_id", "id", "CASCADE"),
    ("stock_check_dish", "store", "store_id", "id", "RESTRICT"),
    ("stock_check_dish", "dish", "dish_id", "id", "RESTRICT"),
    ("stock_check_ingredient", "store", "store_id", "id", "RESTRICT"),
    ("stock_check_ingredient", "ingredient", "ingredient_id", "id", "RESTRICT"),
    ("printer", "store", "store_id", "id", "CASCADE"),
]

# 章节划分
chapter_titles = {
    "sys_role": "一、系统权限与账号",
    "store": "二、门店与门店设置",
    "member_category": "三、客户 / 会员",
    "dish_category": "四、菜品与商品",
    "table_type": "五、台桌、排队、预约",
    "cart": "六、订单",
    "coupon": "七、促销（优惠券）",
    "dish_stock": "八、库存",
    "printer": "九、硬件与意见反馈",
}
chapters = {}
current = None
for t in tables:
    if t["name"] in chapter_titles:
        current = chapter_titles[t["name"]]
        chapters.setdefault(current, [])
    if current:
        chapters[current].append(t)

VIEWS = [
    ("v_dish_flow", "菜品流水", "订单明细级流水记录，支撑商户端【报表-菜品流水】"),
    ("v_receipt_report", "收款日报", "按门店按天汇总实收，支撑【报表-收款报表】"),
    ("v_dish_sales_ranking", "菜品销售排行", "按菜品汇总销量与销售额，支撑【报表-菜品销售排行】"),
    ("v_store_sales_ranking", "门店销售排行", "按门店汇总单量与销售额，支撑【报表-门店销售排行】"),
    ("v_today_income", "今日收入", "收银端/首页今日订单数与实收汇总"),
    ("v_refund_record", "退单记录", "退单明细视图，支撑【报表-退单记录】"),
]


def field_rows(t):
    lines = []
    for c in t["columns"]:
        key = "PK" if c["is_pk"] else ("UK" if c["is_unique"] else "")
        if c["is_auto"]:
            key = (key + "+AI").strip("+") if key else "AI"
        lines.append(f"| `{c['name']}` | {c['type']} | {'否' if not c['nullable'] else '是'} | "
                     f"{c['default'] or '—'} | {key} | {c['comment']} |")
    return lines


# ============================ Markdown ============================
md = []
md.append("# 智慧餐厅 数据库字典\n")
md.append("> 数据库：`smart_restaurant` ｜ MySQL 8.0 ｜ InnoDB ｜ `utf8mb4_general_ci` ｜ 43 张表 + 6 个报表视图\n")
md.append("> 关联完整性：逻辑外键 + 索引保障（不建物理外键约束），ER 关系见【第二部分】\n")
md.append("## 目录\n")
md.append("1. [第一部分 数据表基础说明](#1)")
md.append("2. [第二部分 表间关联关系（ER）](#2)")
md.append("3. [第三部分 字段详细说明](#3)")
md.append("4. [附录 报表视图](#4)")
md.append("")
md.append("## 设计约定\n")
md.append("- 主键统一为 `id BIGINT UNSIGNED AUTO_INCREMENT`；软删除统一使用 `is_deleted`（0否/1是）；时间统一为 `create_time`/`update_time`。")
md.append("- 金额统一 `DECIMAL(10,2)`；状态字段统一 `TINYINT` 并在注释中标注枚举含义；全部表使用 `utf8mb4` + `utf8mb4_general_ci`。")
md.append("- 表间关联采用**逻辑外键**（同名列 + 索引），由应用层保证一致性；本设计不创建物理外键约束，便于水平扩展与批量数据导入。")
md.append("- 并发控制：`orders.version` 乐观锁；`payment_record.idempotency_key` 支付幂等键。\n")

# ---- 第一部分：表基础说明 ----
md.append("## 第一部分 数据表基础说明\n")
md.append("| 序号 | 表名 | 中文释义 | 业务用途 |")
md.append("|---|---|---|---|")
for i, t in enumerate(tables, 1):
    md.append(f"| {i} | `{t['name']}` | {t['comment']} | {TABLE_INFO.get(t['name'], '')} |")
md.append("")

# ---- 第二部分：ER 关联关系 ----
md.append("## 第二部分 表间关联关系（ER）\n")
md.append("关联类型约定：`1:N`＝父表 1 条对应子表 N 条；`N:M`＝多对多（由中间表承载）。级联策略：`CASCADE`＝删除父表级联删除子表；`RESTRICT`＝有子记录时禁止删除父记录；`LOGICAL`＝仅逻辑关联（无级联约束）。\n")
md.append("| 子表 | 父表 | 关联字段(子.字段 = 父.字段) | 关联类型 | 级联策略 |")
md.append("|---|---|---|---|---|")
mid_tables = {"sys_role_permission", "dish_ingredient_rel"}
for sub, par, col, refcol, policy in ER_DATA:
    typ = "N:M(中间表)" if sub in mid_tables else "N:1"
    md.append(f"| `{sub}` | `{par}` | `{sub}.{col}` = `{par}.{refcol}` | {typ} | {policy} |")
md.append("")

# 父表→子表视角概览
md.append("### 父表 → 子表（一对多）速查\n")
parent_map = {}
for sub, par, col, refcol, policy in ER_DATA:
    parent_map.setdefault(par, []).append(sub)
for par, subs in parent_map.items():
    md.append(f"- **{par}** → {', '.join(subs)}")
md.append("")

# ---- 第三部分：字段详细说明 ----
md.append("## 第三部分 字段详细说明\n")
for title, tbls in chapters.items():
    md.append(f"## {title}\n")
    for t in tbls:
        md.append(f"### {t['name']} — {t['comment']}\n")
        md.append(f"业务用途：{TABLE_INFO.get(t['name'], '')}\n")
        md.append("| 字段名 | 数据类型 | 允许空 | 默认值 | 键 | 中文含义 |")
        md.append("|---|---|---|---|---|---|")
        md.extend(field_rows(t))
        md.append("")
        if t["indexes"]:
            md.append("**索引：**")
            for idx in t["indexes"]:
                md.append(f"- `{idx['type']}` `{idx['name']}` ({', '.join('`'+c+'`' for c in idx['cols'])})"
                          + ("（唯一）" if idx["unique"] else ""))
            md.append("")
        # 逻辑外键提示
        rels = [r for r in ER_DATA if r[0] == t["name"]]
        if rels:
            md.append("**逻辑外键：**")
            for r in rels:
                md.append(f"- `{r[0]}.{r[2]}` → `{r[1]}.{r[3]}`（{r[4]}）")
            md.append("")
    md.append("")

# ---- 附录 ----
md.append("## 附录 报表视图\n")
md.append("| 视图名 | 名称 | 用途 |")
md.append("|---|---|---|")
for v in VIEWS:
    md.append(f"| `{v[0]}` | {v[1]} | {v[2]} |")
md.append("")

with open(MD_PATH, "w", encoding="utf-8") as f:
    f.write("\n".join(md))
print(f"[OK] 数据库字典 Markdown 已生成：{MD_PATH}")

# ============================ Word ============================
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def set_cn_font(run, name="宋体", size=None, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        set_cn_font(r, "黑体", 16 if level == 1 else 13, bold=True)
    return h


def add_table(headers, rows, widths=None, font=9):
    tb = doc.add_table(rows=1, cols=len(headers))
    tb.style = "Table Grid"
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htext in enumerate(headers):
        cell = tb.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(htext)
        set_cn_font(r, "黑体", font, bold=True)
    for row_vals in rows:
        row = tb.add_row().cells
        for i, v in enumerate(row_vals):
            row[i].text = ""
            r = row[i].paragraphs[0].add_run(str(v))
            set_cn_font(r, "宋体", font)
    if widths:
        for row in tb.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = w
    return tb


# 封面
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("智慧餐厅 数据库字典")
set_cn_font(r, "黑体", 26, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("数据库 smart_restaurant ｜ MySQL 8.0 ｜ InnoDB ｜ utf8mb4_general_ci ｜ 43 张表 + 6 视图")
set_cn_font(r, "宋体", 12, color=RGBColor(0x66, 0x66, 0x66))
doc.add_paragraph()

add_heading("设计约定", 1)
for line in [
    "1. 主键统一为 id BIGINT UNSIGNED AUTO_INCREMENT；软删除统一使用 is_deleted；时间统一为 create_time / update_time。",
    "2. 金额统一 DECIMAL(10,2)；状态字段统一 TINYINT；全部表使用 utf8mb4 + utf8mb4_general_ci。",
    "3. 表间关联采用逻辑外键（同名列 + 索引），由应用层保证一致性，不创建物理外键约束，便于扩展与批量导入。",
    "4. 并发控制：orders.version 乐观锁；payment_record.idempotency_key 支付幂等键。",
]:
    p = doc.add_paragraph(line)
    set_cn_font(p.runs[0], "宋体", 10.5)
doc.add_page_break()

# 第一部分：表基础说明
add_heading("第一部分 数据表基础说明", 1)
add_table(
    ["序号", "表名", "中文释义", "业务用途"],
    [(i, t["name"], t["comment"], TABLE_INFO.get(t["name"], "")) for i, t in enumerate(tables, 1)],
    [Cm(1.2), Cm(3.0), Cm(4.2), Cm(8.0)],
)
doc.add_page_break()

# 第二部分：ER
add_heading("第二部分 表间关联关系（ER）", 1)
p = doc.add_paragraph("关联类型约定：1:N＝父表 1 条对应子表 N 条；N:M＝多对多（由中间表承载）。级联策略：CASCADE＝级联删除；RESTRICT＝限制删除；LOGICAL＝仅逻辑关联。")
set_cn_font(p.runs[0], "宋体", 10)
add_table(
    ["子表", "父表", "关联字段（子.字段 = 父.字段）", "关联类型", "级联策略"],
    [(s, p_, f"{s}.{c} = {p_}.{rc}", "N:M(中间表)" if s in mid_tables else "N:1", pl)
     for s, p_, c, rc, pl in ER_DATA],
    [Cm(3.2), Cm(3.0), Cm(5.2), Cm(2.6), Cm(2.4)],
)
doc.add_page_break()

# 第三部分：字段详细说明
add_heading("第三部分 字段详细说明", 1)
for title, tbls in chapters.items():
    add_heading(title, 2)
    for t in tbls:
        add_heading(f"{t['name']}  {t['comment']}", 3)
        p = doc.add_paragraph(f"业务用途：{TABLE_INFO.get(t['name'], '')}")
        set_cn_font(p.runs[0], "宋体", 10)
        rows = []
        for c in t["columns"]:
            key = "PK" if c["is_pk"] else ("UK" if c["is_unique"] else "")
            if c["is_auto"]:
                key = (key + "+自增").lstrip("+") if key else "自增"
            rows.append((c["name"], c["type"], "否" if not c["nullable"] else "是",
                         c["default"] or "—", key, c["comment"]))
        add_table(["字段名", "数据类型", "允许空", "默认值", "键", "中文含义"], rows,
                  [Cm(2.6), Cm(3.0), Cm(1.2), Cm(2.6), Cm(1.4), Cm(6.0)])
        if t["indexes"]:
            p = doc.add_paragraph()
            r = p.add_run("索引：")
            set_cn_font(r, "黑体", 9.5, bold=True)
            for idx in t["indexes"]:
                cols = ", ".join(idx["cols"])
                uniq = "（唯一）" if idx["unique"] else ""
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.5)
                r = p.add_run(f"{idx['type']} {idx['name'] or ''} ({cols}) {uniq}".strip())
                set_cn_font(r, "宋体", 9)
        rels = [r for r in ER_DATA if r[0] == t["name"]]
        if rels:
            p = doc.add_paragraph()
            r = p.add_run("逻辑外键：")
            set_cn_font(r, "黑体", 9.5, bold=True)
            for r_ in rels:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.5)
                r = p.add_run(f"{r_[0]}.{r_[2]} → {r_[1]}.{r_[3]}（{r_[4]}）")
                set_cn_font(r, "宋体", 9)
        doc.add_paragraph()

# 附录
add_heading("附录 报表视图", 1)
add_table(["视图名", "名称", "用途"], [(v[0], v[1], v[2]) for v in VIEWS],
          [Cm(3.2), Cm(3.0), Cm(10.0)])

doc.save(DOCX_PATH)
print(f"[OK] 数据库字典 Word 已生成：{DOCX_PATH}")
print(f"     共 {len(tables)} 张表，{sum(len(t['columns']) for t in tables)} 个字段，{len(ER_DATA)} 条逻辑外键关系")
